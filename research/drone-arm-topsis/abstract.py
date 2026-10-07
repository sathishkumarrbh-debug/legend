from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
d=Document(); s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
for a in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,a,Cm(2.54))
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
def para(text='',size=12,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.CENTER,after=6):
    p=d.add_paragraph(); p.alignment=align; p.paragraph_format.space_after=Pt(after)
    if text: r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.italic=italic
    return p
para('International Conference on Advances in Aerospace and Navigation Systems (ICAAN–2026)',11,italic=True,after=0)
para('Track: Propulsion and Structures',11,italic=True,after=14)
para('Material Selection for a Quadcopter Drone Arm Using Beam Theory and TOPSIS Method',14,True,after=12)
p=para(after=2)
auth=[('Dr. S. R. Sathishkumar','1*'),('[Staff Name]','1'),('[Student 1]','2'),('[Student 2]','2'),('[Student 3]','2'),('[Student 4]','2')]
for i,(n,sup) in enumerate(auth):
    r=p.add_run(n); r.bold=True; r=p.add_run(sup); r.font.superscript=True
    if i<len(auth)-1: p.add_run(', ')
para('¹Assistant Professor, ²UG Student, Department of Mechanical Engineering,',11,italic=True,after=0)
para('St. Joseph’s College of Engineering and Technology, Thanjavur – 613 403, Tamil Nadu, India',11,italic=True,after=0)
para('*Corresponding author: [email address]',11,italic=True,after=16)
p=para(align=WD_ALIGN_PARAGRAPH.LEFT); r=p.add_run('Abstract'); r.bold=True
A=('The arm of a multirotor drone must be light to increase flight time, yet stiff and strong enough to carry the motor thrust without excessive bending. This work presents a simple and systematic method to select the best material for the arm of a 450 mm-class quadcopter. The arm was modelled as a cantilever hollow circular tube (12 mm outer diameter, 10 mm inner diameter, 200 mm length) with a design end load of 20 N, obtained from a maximum motor thrust of 10 N and a safety factor of 2. Using classical beam theory, the maximum bending stress (45.5 MPa), tip deflection, arm mass and factor of safety were calculated for five candidate materials: AISI 4130 steel, aluminium 6061-T6, titanium Ti-6Al-4V, glass fibre reinforced polymer (GFRP) and carbon fibre reinforced polymer (CFRP). The materials were then ranked using the TOPSIS multi-criteria decision-making method with mass, deflection, factor of safety and cost as criteria, weighted 0.35, 0.25, 0.20 and 0.20 respectively. CFRP obtained the highest closeness score (0.784), giving the lightest arm (11.1 g, about 80% lighter than steel) with a tip deflection of 1.45 mm and a factor of safety of 13.2. Aluminium 6061-T6 ranked second (0.713) and became the best choice when cost was given the highest weight, which agrees with the use of carbon-fibre arms in premium drones and aluminium arms in low-cost drones. The proposed spreadsheet-based approach is transparent, easy to verify and can be extended to other drone structural components.')
p=para(A,align=WD_ALIGN_PARAGRAPH.JUSTIFY,after=10); p.paragraph_format.line_spacing=1.15
p=para(align=WD_ALIGN_PARAGRAPH.JUSTIFY,after=0); r=p.add_run('Keywords: '); r.bold=True
p.add_run('Quadcopter; Drone arm; Material selection; Cantilever beam; TOPSIS; Carbon fibre composite')
d.save('ICAAN2026_Abstract_DroneArm.docx'); print(len(A.split()),'words')
