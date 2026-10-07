from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
wb=Workbook(); ws=wb.active; ws.title='Beam Calculation'
FN='Times New Roman'
def f(**k): return Font(name=FN,**k)
BLUE=f(color='0000FF'); BOLD=f(bold=True); GREEN=f(color='008000'); NORM=f()
HDR=PatternFill('solid',fgColor='4A1D8C'); HF=f(bold=True,color='FFFFFF'); KEY=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='999999'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,row,cols):
    for i,t in enumerate(cols):
        c=ws.cell(row=row,column=i+1,value=t); c.font=HF; c.fill=HDR; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
ws['A1']='Quadcopter Drone Arm – Cantilever Beam Calculation'; ws['A1'].font=f(bold=True,size=14,color='2A0D5C')
ws['A2']='Blue = input (you may change) · Black = formula · Arm modelled as a hollow circular tube fixed at the frame, motor thrust at the free end'; ws['A2'].font=f(italic=True,size=10)
ws['A3']='INPUTS'; ws['A3'].font=BOLD
inp=[('Outer diameter D (mm)',12),('Inner diameter d (mm)',10),('Arm length L (mm)',200),('Motor maximum thrust (N)',10),('Safety factor on load',2)]
for i,(t,v) in enumerate(inp):
    ws.cell(row=4+i,column=1,value=t).font=NORM; c=ws.cell(row=4+i,column=2,value=v); c.font=BLUE; c.fill=KEY
ws['C7']='≈ 1 kg-f per motor (450 mm class, ~1.5 kg drone)'; ws['C7'].font=f(italic=True,size=10)
ws['A9']='Design load F (N)'; ws['B9']='=B7*B8'
ws['A11']='SECTION PROPERTIES & STRESS'; ws['A11'].font=BOLD
calc=[('Area A = π(D²−d²)/4  (mm²)','=PI()*(B4^2-B5^2)/4'),('Moment of inertia I = π(D⁴−d⁴)/64  (mm⁴)','=PI()*(B4^4-B5^4)/64'),('Section modulus Z = I/(D/2)  (mm³)','=B13/(B4/2)'),('Max bending moment M = F×L  (N·mm)','=B9*B6'),('Max bending stress σ = M/Z  (MPa)','=B15/B14')]
for i,(t,v) in enumerate(calc):
    ws.cell(row=12+i,column=1,value=t).font=NORM; c=ws.cell(row=12+i,column=2,value=v); c.number_format='0.00'
ws['B9'].number_format='0.00'
ws['A18']='MATERIALS'; ws['A18'].font=BOLD
hdr(ws,19,['Material','Density ρ (kg/m³)','Young’s modulus E (GPa)','Yield / tensile strength (MPa)','Approx. cost (₹/kg)','Arm mass m = ρAL (g)','Tip deflection δ = FL³/3EI (mm)','Factor of safety = Strength/σ','Material cost per arm (₹)'])
mats=[('Steel AISI 4130',7850,205,460,250),('Aluminium 6061-T6',2700,69,276,350),('Titanium Ti-6Al-4V',4430,114,880,3500),('GFRP (glass/epoxy)',1900,25,350,700),('CFRP (carbon/epoxy)',1600,70,600,3000)]
for i,m in enumerate(mats):
    r=20+i
    ws.cell(row=r,column=1,value=m[0]).font=NORM
    for j in range(1,5):
        c=ws.cell(row=r,column=j+1,value=m[j]); c.font=BLUE
    ws.cell(row=r,column=6,value=f'=B{r}*$B$12*$B$6*1E-6').number_format='0.00'
    ws.cell(row=r,column=7,value=f'=$B$9*$B$6^3/(3*C{r}*1000*$B$13)').number_format='0.00'
    ws.cell(row=r,column=8,value=f'=D{r}/$B$16').number_format='0.00'
    ws.cell(row=r,column=9,value=f'=F{r}/1000*E{r}').number_format='0.0'
    for j in range(1,10): ws.cell(row=r,column=j).border=BOX
ws['A26']='Notes:'; ws['A26'].font=BOLD
notes=['Material properties: typical handbook values (ASM Handbook / MatWeb). Composites: typical values for tubes, actual values depend on fibre lay-up.',
 'Costs: approximate raw-material prices in India (2026) – only used for comparison; update with local supplier quotes if available.',
 'Strength used: yield strength for metals, tensile strength for composites (composites have no yield point).',
 'Formulas: Strength of Materials – cantilever with end load: M = F·L, σ = M/Z, δ = F·L³/(3·E·I).']
for i,t in enumerate(notes): ws.cell(row=27+i,column=1,value=t).font=f(size=10,italic=True)
ws.column_dimensions['A'].width=44
for col,w in zip('BCDEFGHI',[14,14,15,13,14,16,15,14]): ws.column_dimensions[col].width=w
ws.row_dimensions[19].height=48
for row in ws.iter_rows(min_row=1,max_row=31):
    for c in row:
        if c.font.name!=FN: c.font=f(bold=c.font.bold,italic=c.font.italic,color=c.font.color,size=c.font.size)

# ---- TOPSIS ----
t=wb.create_sheet('TOPSIS')
t['A1']='TOPSIS Ranking of Materials'; t['A1'].font=f(bold=True,size=14,color='2A0D5C')
t['A2']='TOPSIS = Technique for Order of Preference by Similarity to Ideal Solution. Best material = closest to the ideal best and farthest from the ideal worst.'; t['A2'].font=f(italic=True,size=10)
hdr(t,4,['Criterion','Mass (g)','Deflection (mm)','Factor of safety','Cost (₹)','Sum of weights'])
t['A5']='Impact'; t['B5']='Lower is better'; t['C5']='Lower is better'; t['D5']='Higher is better'; t['E5']='Lower is better'
t['A6']='Weight (importance)'
for col,w in zip('BCDE',[0.35,0.25,0.20,0.20]):
    c=t[f'{col}6']; c.value=w; c.font=BLUE; c.fill=KEY; c.number_format='0.00'
t['F6']='=SUM(B6:E6)'; t['F6'].number_format='0.00'
t['A7']='Weights chosen: mass most important for flight time; weights must add up to 1.'; t['A7'].font=f(italic=True,size=10)
t['A8']='Step 1: Decision matrix (from Beam Calculation sheet)'; t['A8'].font=BOLD
for i in range(5):
    r=9+i; s=20+i
    t[f'A{r}']=f"='Beam Calculation'!A{s}"; t[f'A{r}'].font=GREEN
    for col,src in zip('BCDE','FGHI'):
        c=t[f'{col}{r}']; c.value=f"='Beam Calculation'!{src}{s}"; c.font=GREEN; c.number_format='0.00'
t['A14']='√(sum of squares)'
for col in 'BCDE': t[f'{col}14']=f'=SQRT(SUMSQ({col}9:{col}13))'; t[f'{col}14'].number_format='0.000'
t['A16']='Step 2–3: Normalised × weight  (value ÷ √sum of squares × weight)'; t['A16'].font=BOLD
for i in range(5):
    r=17+i
    t[f'A{r}']=f'=A{9+i}'
    for col in 'BCDE': t[f'{col}{r}']=f'={col}{9+i}/{col}$14*{col}$6'; t[f'{col}{r}'].number_format='0.0000'
t['A22']='Step 4: Ideal best (V+)'; t['A23']='Ideal worst (V−)'
for col,best in zip('BCDE',['MIN','MIN','MAX','MIN']):
    worst='MAX' if best=='MIN' else 'MIN'
    t[f'{col}22']=f'={best}({col}17:{col}21)'; t[f'{col}23']=f'={worst}({col}17:{col}21)'
    t[f'{col}22'].number_format=t[f'{col}23'].number_format='0.0000'
t['A25']='Step 5–6: Distances, closeness score and rank'; t['A25'].font=BOLD
hdr(t,26,['Material','D+ (distance to best)','D− (distance to worst)','Closeness C = D−/(D+ + D−)','Rank'])
for i in range(5):
    r=27+i; v=17+i
    t[f'A{r}']=f'=A{9+i}'
    t[f'B{r}']=f'=SQRT((B{v}-B$22)^2+(C{v}-C$22)^2+(D{v}-D$22)^2+(E{v}-E$22)^2)'
    t[f'C{r}']=f'=SQRT((B{v}-B$23)^2+(C{v}-C$23)^2+(D{v}-D$23)^2+(E{v}-E$23)^2)'
    t[f'D{r}']=f'=C{r}/(B{r}+C{r})'
    t[f'E{r}']=f'=RANK(D{r},$D$27:$D$31,0)'
    for col in 'BCD': t[f'{col}{r}'].number_format='0.000'
    for col in 'ABCDE': t[f'{col}{r}'].border=BOX
t['A33']='Best material:'; t['A33'].font=BOLD
t['B33']='=INDEX(A27:A31,MATCH(1,E27:E31,0))'; t['B33'].font=f(bold=True,color='4A1D8C')
t.column_dimensions['A'].width=40
for col in 'BCDEF': t.column_dimensions[col].width=18
t.row_dimensions[4].height=32; t.row_dimensions[26].height=40
for row in t.iter_rows(min_row=1,max_row=33):
    for c in row:
        if c.font.name!=FN: c.font=f(bold=c.font.bold,italic=c.font.italic,color=c.font.color,size=c.font.size)
wb.save('Drone_Arm_Calculation.xlsx')
